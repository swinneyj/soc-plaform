"""Generate a richer SOC Platform architecture PPTX.

This script builds a presentation that you could actually walk through
in a design review: it includes high-level architecture, per-tab UI
mapping, API/DB/tool responsibilities, extension patterns, and runtime
operations.

Usage:
    pip install python-pptx
    python generate_architecture_pptx.py

The output file will be created as:
    SOC_Automation_Architecture.pptx
"""

from __future__ import annotations

import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


ROOT = Path(__file__).resolve().parent


def rel(path: Path) -> str:
    """Return a workspace-relative, POSIX-style path string."""
    try:
        return path.relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


# Simple visual theme (dark background, light text, blue accent)
BG_COLOR = RGBColor(14, 22, 40)        # deep navy
TITLE_COLOR = RGBColor(235, 241, 255)  # near-white
BODY_COLOR = RGBColor(204, 214, 228)   # light gray-blue
ACCENT_COLOR = RGBColor(80, 156, 255)  # blue accent


def set_slide_background(slide) -> None:
    """Apply a consistent dark background to a slide."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def style_title(shape) -> None:
    """Style the slide title with accent color and larger font."""
    if not shape or not shape.text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    font = p.font
    font.size = Pt(40)
    font.bold = True
    font.color.rgb = ACCENT_COLOR


def style_body_frame(tf) -> None:
    """Style all paragraphs in a body text frame."""
    for i, p in enumerate(tf.paragraphs):
        font = p.font
        font.size = Pt(20 if i == 0 else 18)
        font.color.rgb = BODY_COLOR
        if i == 0:
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.LEFT


def add_title_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[0]  # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    set_slide_background(slide)

    title.text = "SOC Platform Architecture Overview"
    subtitle.text = (
        "Local SOC orchestration platform – web UI, FastAPI, "
        "PostgreSQL, Redis, and Ollama"
    )

    style_title(title)
    # Style subtitle
    if subtitle and subtitle.text_frame.paragraphs:
        p = subtitle.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        font = p.font
        font.size = Pt(18)
        font.color.rgb = BODY_COLOR


def add_agenda_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]  # Title + content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    title.text = "Agenda"
    body.text = "What this deck covers"

    for line in [
        "High-level architecture and runtime services",
        "Dashboard tabs mapped to files and endpoints",
        "API, database, and tools responsibilities",
        "AI/Ollama integration for analysis and code review",
        "Safe extension patterns (UI, API, tools, DB)",
        "Operations model: containers, handoff, and runbooks",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_high_level_architecture_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    api_main = ROOT / "api" / "main.py"
    ui_index = ROOT / "web" / "index.html"
    db_models = ROOT / "db" / "models.py"
    docker_compose = ROOT / "docker-compose.yml"

    title.text = "High-Level Architecture"
    body.text = "Layered local SOC platform built to be portable and operator-friendly."

    for line in [
        f"Web UI: {rel(ui_index)} (Vue-style single page, TailwindCSS).",
        f"API: {rel(api_main)} (FastAPI on uvicorn, port 8000).",
        f"Database: {rel(db_models)} (SQLAlchemy models – Postgres runtime, SQLite fallback).",
        "Async jobs: Redis-backed queue for long-running tool executions.",
        "AI layer: Ollama service on host for local LLM inference.",
        f"Containers: {rel(docker_compose)} orchestrates api-service, postgres, and redis.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_ui_tabs_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    ui_path = ROOT / "web" / "index.html"

    title.text = "Dashboard Tabs (web/index.html)"
    body.text = f"Entry point: {rel(ui_path)}"

    for line in [
        "Tools – catalog from Commander_Registry.json, run scripts in Tools/.",
        "Database – triage cases, pasted notables, verdict breakdowns.",
        "AI Analysis – case picker, model selector, supportive queries.",
        "Closure Notes – rule-driven closure templates and generated notes.",
        "Jobs – history of tool executions and logs.",
        "Reports – download generated reports/briefings.",
        "Code Review – section detection and function-level LLM reviews.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_code_review_flow_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    api_main = ROOT / "api" / "main.py"
    ui_index = ROOT / "web" / "index.html"

    title.text = "Code Review Flow"
    body.text = "How the Code Review tab uses API + Ollama to review functions."

    for line in [
        f"UI: Code Review tab in {rel(ui_index)} loads code via file/folder upload.",
        "Detect Sections – calls /api/code-review/sections to extract functions/methods.",
        "Section focus – clicking a section narrows the textarea to that function body.",
        "Submit for Review – posts to /api/code-review with snippet + language/model.",
        f"API logic: implemented in {rel(api_main)}, stores results in CodeReview table.",
        "AI: API uses Ollama service to generate recommendations for that specific function.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_api_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    api_main = ROOT / "api" / "main.py"

    title.text = "API Layer (FastAPI)"
    body.text = f"Entry point: {rel(api_main)}"

    for line in [
        "Tools: /api/tools, /api/tools/{name}, /api/execute, /api/jobs/{id}.",
        "Database: /api/db/stats, /api/db/triage, /api/db/triage/{case_id}.",
        "Pasted notables: ingest, list, promote to triage, delete workflows.",
        "Supportive queries: CRUD for per-rule SPL and storage of results.",
        "Code review: submit snippets, store results, and detect sections.",
        "Health: /health and /api/health endpoints for basic monitoring.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_db_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    db_models = ROOT / "db" / "models.py"

    title.text = "Database Layer (SQLAlchemy)"
    body.text = f"Models file: {rel(db_models)}"

    for line in [
        "TriageResult – per-case verdict, confidence, summary, remediation steps.",
        "SplunkEvent – raw event rows backing case context.",
        "AnalysisResult – Ollama analysis output tied to a case + model.",
        "ESCorrelationRule – rule metadata, closure requirements, severity.",
        "ClosureNote – generated incident closure notes lifecycle.",
        "SupportiveQuery & SupportiveQueryResult – enrichment SPL and stored evidence.",
        "PlaceholderAlias – maps logical placeholders (host/dest/user) to fields.",
        "CodeReview – stores LLM review results for code snippets.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_tools_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    tools_root = ROOT / "Tools"
    tool_indexer = tools_root / "tool_indexer" / "tool_indexer.py"
    platform_blueprint = tools_root / "platform_blueprint" / "platform_blueprint.py"
    commander_registry = ROOT / "Commander_Registry.json"

    title.text = "Tools & Orchestration"
    body.text = f"Tools root: {rel(tools_root)}"

    for line in [
        f"Commander registry: {rel(commander_registry)} – indexed tool metadata.",
        f"Tool indexer: {rel(tool_indexer)} – parses TOOL_NAME/DESC/CATEGORY/ARG headers.",
        f"Platform blueprint: {rel(platform_blueprint)} – exports SOC_Architecture_Blueprint.md.",
        "Each tool lives under Tools/<category>/ and is just a script with headers.",
        "UI reads Commander_Registry.json to show a searchable, filterable tool catalog.",
        "Common categories: ingestion pipelines, intel, health checks, maintenance, utilities.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_ai_services_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    ollama_service = ROOT / "services" / "ollama_service.py"

    title.text = "AI / Ollama Integration"
    body.text = f"Client wrapper: {rel(ollama_service)}"

    for line in [
        "OllamaClient checks /api/tags and exposes list_models() and generate().",
        "Default model is configurable (e.g., llama3.1:8b, qwen variants).",
        "analyze_security_event() provides a SOC-tuned prompt template.",
        "AI Analysis tab calls API → Ollama → stores AnalysisResult.",
        "Code Review endpoints also rely on the same Ollama service.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_extension_patterns_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    title.text = "Safe Extension Patterns"
    body.text = "How to extend UI, API, tools, and DB without surprises."

    for line in [
        "UI: add a new tab and state in web/index.html; call new /api/... endpoints, mirror existing patterns.",
        "API: add FastAPI routes in api/main.py, use Pydantic models, and keep DB access in try/finally blocks.",
        "Tools: create Tools/<category>/<tool_name>.py with TOOL_NAME/DESC/CATEGORY headers, then run tool_indexer.",
        "DB: extend db/models.py with new tables, let SQLAlchemy create schema, and wire into new endpoints.",
        "Config: prefer environment variables (DATABASE_URL, OLLAMA_URL, ports) over hard-coded paths.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def add_operations_slide(prs: Presentation) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1].text_frame

    set_slide_background(slide)

    project_spec = ROOT / "PROJECT_SPEC.md"
    runbook = ROOT / "RUNBOOK.md"
    cheat_sheet = ROOT / "OPERATOR_CHEAT_SHEET.md"
    docker_compose = ROOT / "docker-compose.yml"

    title.text = "Operations & Runtime Model"
    body.text = "How operators run, share, and troubleshoot the platform."

    for line in [
        f"Project spec: {rel(project_spec)} – high-level architecture and flows.",
        f"Runbook: {rel(runbook)} – bootstrap, DB handoff, and recommended team model.",
        f"Operator cheat sheet: {rel(cheat_sheet)} – start/stop/status and safe Git pull.",
        f"Compose stack: {rel(docker_compose)} – api-service, postgres, redis containers.",
        "Shared DB handoff uses a PostgreSQL dump on a shared path, outside Git.",
        "Scripts under scripts/ wrap common flows: start_platform, stop_platform, troubleshoot_platform.",
    ]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    style_title(title)
    style_body_frame(body)


def generate_pptx(output_path: Path | None = None) -> Path:
    """Generate the architecture PPTX and return its path."""
    prs = Presentation()

    add_title_slide(prs)
    add_agenda_slide(prs)
    add_high_level_architecture_slide(prs)
    add_ui_tabs_slide(prs)
    add_code_review_flow_slide(prs)
    add_api_slide(prs)
    add_db_slide(prs)
    add_tools_slide(prs)
    add_ai_services_slide(prs)
    add_extension_patterns_slide(prs)
    add_operations_slide(prs)

    if output_path is None:
        output_path = ROOT / "SOC_Automation_Architecture.pptx"

    prs.save(str(output_path))
    return output_path


if __name__ == "__main__":
    out = generate_pptx()
    print(f"[+] Architecture PPTX generated: {out}")
